from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph


REPLACEMENTS = {
    "电路板正中央为大容量高性能微控制器（MCU），采用四面引脚贴片封装，作为整机运算核心，核心功能如下：": (
        "电路板正中央计划采用高性能微控制器（MCU）作为整机运算核心；具体型号、主频、"
        "Flash/RAM容量、封装规格和关键引脚分配尚未确定，待后续选型后补充。"
    ),
    "1. 输入防护：电池输入端搭载防反接二极管、自恢复保险丝、TVS瞬态抑制管，全面防范电源接反、线路短路、浪涌冲击等问题，避免主板硬件损坏；": (
        "1. 输入防护：计划配置防反接、过流和浪涌保护；保护器件型号、耐压等级、"
        "动作阈值及选型计算尚未完成，待补充。"
    ),
    "2. 多级降压：锂电池高压经一级降压至5V，再通过LDO线性稳压输出纯净3.3V电压，单独供给MCU、IMU、GPS等高精度传感器，保障采样数据稳定；": (
        "2. 多级降压：电源架构暂定提供5V与3.3V两路供电；稳压芯片型号、额定输出电流、"
        "转换效率、纹波指标和热设计参数尚未确定，待补充。"
    ),
    "1. 上电电源测试：分别测试空载、满载工况下3.3V、5V电压纹波参数，验证主板稳压电路工作稳定性；": (
        "1. 上电电源测试：上电后检查3.3V、5V电压是否稳定；负载工况、测量点、"
        "仪器配置、纹波上限和通过标准待补充。"
    ),
    "2. 传感器校准测试：在常温、高低温极限环境下校准陀螺仪、加速度计，检测传感器数据干扰漂移值，保障姿态采集精度；": (
        "2. 传感器校准测试：检查陀螺仪、加速度计数据是否正常；测试温区、采样时长、"
        "零偏与漂移阈值待补充。"
    ),
    "3. EMC干扰测试：在电机满负荷运转工况下采集IMU原始数据，验证地分割、滤波电路的抗干扰性能；": (
        "3. EMC干扰测试：观察电机运行时IMU数据是否受到干扰；干扰源配置、负载工况、"
        "采样方法和判定阈值待补充。"
    ),
    "4. 接口可靠性测试：对各拓展端子进行反复插拔测试，检测通信数据无丢包、无断连问题；": (
        "4. 接口可靠性测试：检查各扩展端子的通信和插拔可靠性；插拔次数、通信速率、"
        "测试时长和允许丢包率待补充。"
    ),
    "5. 整机装机试飞：完成整机装配后，开展悬停、定点、高速机动飞行测试，全面验证飞控控制逻辑与硬件整体稳定性。": (
        "5. 整机装机试飞：计划开展悬停、定点和机动飞行测试；测试科目、风险边界、"
        "记录项目和通过条件待补充。"
    ),
}


def _replace_text(paragraph: Paragraph, replacement: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = replacement
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(replacement)


def build(source: Path, output: Path) -> None:
    if output.exists():
        raise SystemExit(f"refusing to overwrite existing file: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source, output)

    document = Document(output)
    replaced: set[str] = set()
    for paragraph in document.paragraphs:
        if paragraph.text in REPLACEMENTS:
            original = paragraph.text
            _replace_text(paragraph, REPLACEMENTS[original])
            replaced.add(original)

    missing = set(REPLACEMENTS) - replaced
    if missing:
        output.unlink(missing_ok=True)
        raise SystemExit(f"source document changed; replacements not found: {sorted(missing)}")

    document.core_properties.title = "无人机硬件研发文档（首审待修改版）"
    document.core_properties.subject = "审稿 Fixture：关键选型参数和验证标准待补充"
    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description="Create the first-round drone document that needs revision")
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    if not args.source.is_file():
        raise SystemExit(f"DOCX not found: {args.source}")
    build(args.source, args.output)
    print(args.output.resolve())


if __name__ == "__main__":
    main()
