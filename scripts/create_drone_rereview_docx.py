from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph


REPLACEMENTS = {
    "本款无人机主控硬件板为一体化飞控核心主板，集成姿态解算、动力驱动、外设通信、定位导航全功能电路，是小型多旋翼无人机的核心控制单元。主板采用正方形紧凑型布局，兼顾体积轻量化与电磁兼容性，适配消费级、工业小型巡检无人机装机需求。": (
        "本款无人机主控硬件板为一体化飞控核心主板，集成姿态解算、动力驱动、外设通信与定位导航电路。"
        "当前处于工程样机详细设计阶段：原理图框架、PCB布局布线和接口分区已经形成，核心器件按下述候选方案进入定型。"
        "第六章所列数据均为待验证的设计目标，第七章给出样机验收方法；完成样机测试前，不将目标值表述为实测性能。"
    ),
    "电路板正中央计划采用高性能微控制器（MCU）作为整机运算核心；具体型号、主频、Flash/RAM容量、封装规格和关键引脚分配尚未确定，待后续选型后补充。": (
        "工程样机候选主控为STM32F405VGT6（Cortex-M4F，168MHz，1MB Flash，192KB SRAM，LQFP100）。"
        "候选依据是其具备浮点运算单元、多路SPI/I2C/UART、充足的定时器PWM通道和成熟飞控生态，可覆盖1kHz IMU采样与500Hz控制环的设计目标。"
        "定型条件为台架测试中控制任务CPU占用率不高于60%、RAM保留不少于30%，并连续运行30分钟无控制周期超时。"
    ),
    "芯片四周预留充足铺铜与过孔，增强散热，同时隔离数字高速信号线，降低噪声干扰。": (
        "初步资源分配为：SPI1连接IMU，UART1连接GPS，UART2连接433MHz数传，I2C1用于扩展传感器，"
        "TIM1/TIM8提供六路PWM资源；当前PCB布出四路电机输出，并为另外两路保留测试点，六轴应用须在后续板版完成布线与验证。"
        "SWD调试口和关键串口均保留测试点。芯片周围通过铺铜与过孔散热，高速数字线与传感器区域隔离。"
    ),
    "1. 输入防护：计划配置防反接、过流和浪涌保护；保护器件型号、耐压等级、动作阈值及选型计算尚未完成，待补充。": (
        "1. 输入防护：输入范围按2S~4S锂电池的6.0V~16.8V设计。工程样机候选链路为3A自恢复保险丝、"
        "耐压不低于40V的反接保护MOSFET和SMBJ18A TVS。选型以16.8V最高持续输入、20%以上耐压余量及浪涌钳位后不超过后级36V耐压为约束，"
        "样机阶段通过反接、短路和浪涌测试确认器件定型。"
    ),
    "2. 多级降压：电源架构暂定提供5V与3.3V两路供电；稳压芯片型号、额定输出电流、转换效率、纹波指标和热设计参数尚未确定，待补充。": (
        "2. 多级降压：5V负载预算为GPS 0.15A、433MHz数传峰值0.50A、外设0.80A及0.35A余量，设计电流1.80A，"
        "候选TPS5430（36V输入、3A输出）；3.3V负载预算为MCU 0.12A、IMU 0.02A、逻辑与扩展0.36A，设计电流0.50A，"
        "候选TPS62162（由5V母线输入、1A输出），IMU另设低噪声3.3V LDO。验收目标为满载效率不低于85%、5V纹波不高于50mVpp、"
        "数字3.3V纹波不高于30mVpp、IMU电源纹波不高于10mVpp，25℃环境满载30分钟后关键器件表面温度不高于85℃。"
    ),
    "3. 电源隔离：电机驱动电源与传感器电源采用独立供电回路，杜绝电机大电流启停带来的电压波动，保障飞控姿态解算精准稳定。": (
        "3. 电源隔离：电机功率回路不经过传感器供电支路，5V主电源、数字3.3V和IMU低噪声3.3V分别滤波；"
        "功率地与信号地在电源入口附近单点汇接，电机控制信号旁配置连续回流路径。以电机阶跃负载时各电源轨不超出上述纹波和电压容差作为布局定型依据。"
    ),
    "6.1 核心性能": "6.1 设计目标指标（待样机验证）",
    "1. 实时姿态解算刷新率≥500Hz，无人机飞行姿态响应灵敏、操控流畅；": (
        "1. 姿态解算刷新率设计目标≥500Hz：依据STM32F405VGT6 168MHz主频和IMU经SPI1以1kHz采样的候选方案估算；"
        "须通过任务周期日志证明CPU占用率≤60%，连续30分钟无周期超时后方可转为实测指标；"
    ),
    "2. 支持GPS定点、定高、自动返航、航线规划等主流智能飞行模式；": (
        "2. GPS定点、定高、自动返航和航线规划为软件与接口设计目标；本版仅完成GPS/UART接口预留，功能状态为待固件联调和试飞验证；"
    ),
    "3. 搭载四路独立无刷电机调速输出，可适配四轴、六轴多旋翼无人机设备；": (
        "3. 当前PCB提供四路独立无刷电机PWM输出，目标适配四轴无人机；MCU资源预留六路PWM，但六轴适配需增加两路板级输出并完成负载与试飞验证，不作为本版已实现能力；"
    ),
    "4. 预留多串口、SPI、I2C拓展总线，兼容光流、避障、云台、高清图传等各类外接外设；": (
        "4. 多串口、SPI和I2C总线已作为扩展接口预留；光流、避障、云台和图传属于兼容性目标，须分别通过接口电平、带宽和连续通信测试；"
    ),
    "5. 支持宽电压电池输入，可适配2S~4S锂电池供电的小型无人机。": (
        "5. 电源输入设计范围为6.0V~16.8V，对应2S~4S锂电池；须完成低压、标称电压、满充电压和浪涌测试后确认支持范围。"
    ),
    "1. 上电电源测试：上电后检查3.3V、5V电压是否稳定；负载工况、测量点、仪器配置、纹波上限和通过标准待补充。": (
        "1. 上电电源测试：在6.0V、7.4V、11.1V、14.8V和16.8V输入下分别测试空载、50%负载和满载；"
        "使用20MHz带宽限制及短接地弹簧测量5V、数字3.3V和IMU 3.3V测试点。通过条件为稳态电压误差≤±3%，"
        "纹波分别≤50mVpp、30mVpp和10mVpp，满载连续30分钟无复位且关键器件表面温度≤85℃。"
    ),
    "2. 传感器校准测试：检查陀螺仪、加速度计数据是否正常；测试温区、采样时长、零偏与漂移阈值待补充。": (
        "2. 控制周期与传感器测试：在-10℃、25℃和60℃各恒温20分钟后静置采集10分钟IMU数据，验证1kHz采样与500Hz控制环。"
        "通过条件为校准后陀螺仪三轴零偏绝对值≤0.5°/s、加速度模长误差≤0.03g、连续30分钟控制周期超时次数为0，"
        "并记录CPU占用率和RAM峰值以完成MCU定型。"
    ),
    "3. EMC干扰测试：观察电机运行时IMU数据是否受到干扰；干扰源配置、负载工况、采样方法和判定阈值待补充。": (
        "3. EMC干扰测试：四路电机依次在0%、50%、80%占空比和0%→80%阶跃工况运行，每个稳态工况采集5分钟IMU与电源数据。"
        "相对电机停转基线，陀螺仪噪声RMS增幅应≤20%，控制器不得复位，5V与3.3V电源不得超出第一项纹波及电压容差。"
    ),
    "4. 接口可靠性测试：检查各扩展端子的通信和插拔可靠性；插拔次数、通信速率、测试时长和允许丢包率待补充。": (
        "4. 接口可靠性测试：连接器完成100次插拔后，UART以1Mbps、SPI以8MHz、I2C以400kHz分别连续传输30分钟。"
        "通过条件为无断连、无锁死，误码或丢包率≤1×10^-5，接触电阻变化≤20mΩ；所有测试记录固件版本、线缆长度和错误计数。"
    ),
    "5. 整机装机试飞：计划开展悬停、定点和机动飞行测试；测试科目、风险边界、记录项目和通过条件待补充。": (
        "5. 整机装机试飞：先在防护网和系留条件下完成解锁、低油门、悬停、定点、返航和中速机动，再进入空旷场地测试。"
        "单架次悬停不少于10分钟；通过条件为全程无复位或电机失步、姿态跟踪误差RMS≤3°、GPS良好条件下定点水平误差≤1.5m，"
        "并完整保存电源、IMU、控制周期、GPS状态和故障日志。"
    ),
    "|（注：部分内容可能由 AI 生成）": "注：本文部分文字在AI辅助下整理，硬件设计方案、图纸与技术判断由研发人员完成并负责核验。",
}


def _replace_text(paragraph: Paragraph, replacement: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = replacement
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(replacement)


def build(source: Path, output: Path) -> None:
    if output.exists():
        raise SystemExit(f"refusing to overwrite existing file: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source, output)

    document = Document(output)
    replaced: set[str] = set()
    for paragraph in document.paragraphs:
        if paragraph.text in REPLACEMENTS:
            original = paragraph.text
            _replace_text(paragraph, REPLACEMENTS[original])
            replaced.add(original)

    missing = set(REPLACEMENTS) - replaced
    if missing:
        output.unlink(missing_ok=True)
        raise SystemExit(f"source document changed; replacements not found: {sorted(missing)}")

    document.core_properties.title = "无人机硬件研发文档（二次审稿修改版）"
    document.core_properties.subject = "已补充设计阶段、关键器件选型依据和量化验收标准"
    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description="Create the revised drone document for second-round review")
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    if not args.source.is_file():
        raise SystemExit(f"DOCX not found: {args.source}")
    build(args.source, args.output)
    print(args.output.resolve())


if __name__ == "__main__":
    main()
